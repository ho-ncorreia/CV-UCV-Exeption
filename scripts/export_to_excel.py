"""Export the in-memory Sales Credit Rules database to an Excel file.

The database lives in the Node API process, so the API must be running (npm run dev).
Usage: py scripts/export_to_excel.py [--url http://localhost:3001] [--output file.xlsx]
"""

import argparse
import json
import sys
import urllib.error
import urllib.request
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "04246A"
RED = "E30613"

# (field, header, lookup table used to resolve a name column or None)
RULE_COLUMNS = [
    ("id", "ID", None),
    ("status", "Status", None),
    ("type", "Type", None),
    ("blSpecific", "BL Specific", None),
    ("carrier", "Carrier", "carriers"),
    ("effectiveDate", "Effective date", None),
    ("expireDate", "Expire date", None),
    ("requestorId", "Requestor ID", None),
    ("requestorName", "Requestor Name", None),
    ("requestorBranch", "Requestor Sales Branch Code", "branches"),
    ("createDate", "Create date", None),
    ("lastUpdate", "Last update", None),
    ("comments", "Comments", None),
    ("salesCreditBp", "Sales Credit BP", "businessPartners"),
    ("contractHolder", "Contract Holder", "businessPartners"),
    ("contractNumber", "Contract Number", "contracts"),
    ("bookingOffice", "Booking Office", "offices"),
    ("polZone", "Trigram POL Zone", "zones"),
    ("originCountry", "Origin Country", "countries"),
    ("originPoint", "Origin Point", "points"),
    ("podZone", "Trigram POD Zone", "zones"),
    ("destCountry", "Destination Country", "countries"),
    ("destPoint", "Destination Point", "points"),
    ("shipperBp", "Shipper BP Cd", "businessPartners"),
    ("forwarderBp", "Forwarder BP Cd", "businessPartners"),
    ("consigneeBp", "Consignee BP Cd", "businessPartners"),
    ("notifyBp", "Notify Party BP Cd", "businessPartners"),
    ("blNumber", "BL Number", None),
]

DATE_FIELDS = {"effectiveDate", "expireDate"}
DATETIME_FIELDS = {"createDate", "lastUpdate"}

LOOKUP_SHEETS = {
    "types": "Types",
    "carriers": "Carriers",
    "branches": "Sales Branches",
    "businessPartners": "Business Partners",
    "contracts": "Contracts",
    "offices": "Booking Offices",
    "zones": "Zones",
    "countries": "Countries",
    "points": "Points",
}


def fetch_json(base_url: str, path: str):
    with urllib.request.urlopen(f"{base_url.rstrip('/')}{path}", timeout=10) as resp:
        return json.load(resp)


def to_cell_value(field: str, value):
    if not value:
        return None
    if field in DATE_FIELDS:
        return date.fromisoformat(value)
    if field in DATETIME_FIELDS:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
    return value


def style_sheet(ws, date_columns=(), datetime_columns=()):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=NAVY)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.tabColor = RED if ws.title == "Rules" else NAVY

    for idx in date_columns:
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell.number_format = "DD.MM.YYYY"
    for idx in datetime_columns:
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell.number_format = "DD.MM.YYYY HH:MM"

    for col_cells in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(width + 2, 10), 50)


def build_rules_sheet(ws, rules, lookups):
    names = {
        table: {item["code"]: item.get("name", "") for item in items}
        for table, items in lookups.items()
    }

    headers, date_cols, datetime_cols = [], [], []
    for field, header, table in RULE_COLUMNS:
        headers.append(header)
        if field in DATE_FIELDS:
            date_cols.append(len(headers))
        if field in DATETIME_FIELDS:
            datetime_cols.append(len(headers))
        if table:
            headers.append(f"{header} - Name")
    ws.append(headers)

    for rule in sorted(rules, key=lambda r: int(r["id"])):
        row = []
        for field, _, table in RULE_COLUMNS:
            value = rule.get(field, "")
            row.append(to_cell_value(field, value))
            if table:
                row.append(names[table].get(value) or None)
        ws.append(row)

    style_sheet(ws, date_cols, datetime_cols)


def build_lookup_sheet(ws, items):
    columns = list(dict.fromkeys(key for item in items for key in item))
    ws.append([c[0].upper() + c[1:] for c in columns])
    for item in items:
        ws.append([item.get(c) for c in columns])
    style_sheet(ws)


def main():
    parser = argparse.ArgumentParser(description="Export Sales Credit Rules to Excel")
    parser.add_argument("--url", default="http://localhost:3001", help="API base URL")
    parser.add_argument(
        "--output",
        default=f"sales_credit_rules_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
        help="Output .xlsx file",
    )
    args = parser.parse_args()

    try:
        rules = fetch_json(args.url, "/api/rules")
        lookups = fetch_json(args.url, "/api/lookups")
    except urllib.error.URLError as exc:
        sys.exit(f"Cannot reach the API at {args.url} ({exc.reason}). Start it with 'npm run dev'.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"
    build_rules_sheet(ws, rules, lookups)
    for table, title in LOOKUP_SHEETS.items():
        build_lookup_sheet(wb.create_sheet(title), lookups.get(table, []))

    wb.save(args.output)
    print(f"{len(rules)} rule(s) exported to {args.output}")


if __name__ == "__main__":
    main()
